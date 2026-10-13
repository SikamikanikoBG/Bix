"""Meeting minutes as a Word document (.docx): the same content as the e-mail,
laid out for printing and editing. Action items are a table (task, owner, due),
headings use Word's own heading styles, so the navigation pane and a table of
contents work. cite=False leaves out the transcript line numbers."""

from __future__ import annotations

import io
from typing import Any

from bix.server.minutes import L10N, footnote

INK = "111827"
MUTED = "6B7280"
ACCENT = "1F4E79"
HEAD_FILL = "E8EEF5"
LINE = "C9D3DE"


def render_docx(m: dict[str, Any], meta: dict[str, Any], language: str = "bg", cite: bool = True) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.shared import Cm, Pt, RGBColor

    t = L10N.get(language, L10N["en"])
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)  # A4
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.0)

    rgb = lambda hex_: RGBColor.from_string(hex_)  # noqa: E731
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size, normal.font.color.rgb = "Calibri", Pt(11), rgb(INK)
    _east_asian_font(normal, "Calibri")
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, before in (("Title", 22, 0), ("Heading 1", 14, 16)):
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.bold, st.font.color.rgb = "Calibri", Pt(size), True, rgb(ACCENT if name != "Title" else INK)
        _east_asian_font(st, "Calibri")
        st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(before), Pt(4)
    _no_title_border(doc.styles["Title"])

    def refs(p, rs: list[int]) -> None:
        if cite and rs:
            r = p.add_run(f"  ({t['lines']} {', '.join(map(str, rs))})")
            r.font.size, r.font.color.rgb = Pt(9), rgb(MUTED)

    def muted(text: str, size: float = 10, italic: bool = False):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.font.size, r.font.color.rgb, r.italic = Pt(size), rgb(MUTED), italic
        return p

    def bullets(items: list[dict], text_of, bold_of=None) -> None:
        if not items:
            muted(t["none"], 11)
            return
        for it in items:
            p = doc.add_paragraph(style="List Bullet")
            if bold_of:
                p.add_run(bold_of(it)).bold = True
                p.add_run(": ")
            p.add_run(text_of(it))
            refs(p, it.get("refs") or [])

    doc.core_properties.title = meta.get("title") or m.get("title") or ""
    doc.core_properties.author = "Bix"
    doc.add_paragraph(meta.get("title") or m.get("title") or "", style="Title")
    facts = [x for x in (meta.get("date", ""), meta.get("duration", ""), meta.get("room", "")) if x]
    if facts:
        muted(" · ".join(facts))
    if meta.get("participants"):
        p = muted(f"{t['participants']}: ")
        p.add_run(", ".join(meta["participants"])).font.size = Pt(10)

    doc.add_heading(t["summary"], level=1)
    doc.add_paragraph(m.get("summary") or t["none"])

    doc.add_heading(t["decisions"], level=1)
    bullets(m.get("decisions") or [], lambda d: d["text"])

    doc.add_heading(t["actions"], level=1)
    actions = m.get("action_items") or []
    if actions:
        cols = [t["task"], t["owner"], t["due"]] + ([t["lines"].capitalize()] if cite else [])
        widths = [Cm(9.4), Cm(3.6), Cm(2.6)] + ([Cm(1.0)] if cite else [])
        if not cite:
            widths[0] = Cm(10.4)
        table = doc.add_table(rows=1, cols=len(cols))
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        _borders(table)
        for i, label in enumerate(cols):
            cell = table.rows[0].cells[i]
            cell.text = ""
            r = cell.paragraphs[0].add_run(label)
            r.bold, r.font.size = True, Pt(10)
            _shade(cell, HEAD_FILL)
        _repeat_header(table.rows[0])
        for a in actions:
            row = table.add_row().cells
            vals = [a.get("task") or "", a.get("owner") or "", a.get("due") or ""] + ([", ".join(map(str, a.get("refs") or []))] if cite else [])
            for i, v in enumerate(vals):
                row[i].text = ""
                r = row[i].paragraphs[0].add_run(v)
                if i == 1 and v:
                    r.bold = True
                if i == 3:
                    r.font.size, r.font.color.rgb = Pt(9), rgb(MUTED)
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = w
                row.cells[i].paragraphs[0].paragraph_format.space_after = Pt(0)
    else:
        muted(t["none"], 11)

    doc.add_heading(t["questions"], level=1)
    bullets(m.get("open_questions") or [], lambda q: q["text"])

    if m.get("topics"):
        doc.add_heading(t["topics"], level=1)
        bullets(m["topics"], lambda tp: tp["summary"], bold_of=lambda tp: tp["title"])

    doc.add_paragraph()
    muted(footnote(t, cite), 9, italic=True)
    if meta.get("url"):
        _hyperlink(doc.add_paragraph(), meta["url"], t["open"], ACCENT)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --- small OOXML helpers python-docx has no API for ---------------------------------


def _east_asian_font(style, name: str) -> None:
    """Cyrillic text takes the theme font unless rFonts says otherwise for every script."""
    from docx.oxml.ns import qn

    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        fonts.attrib.pop(qn(attr), None)


def _no_title_border(style) -> None:
    from docx.oxml.ns import qn

    ppr = style.element.find(qn("w:pPr"))
    if ppr is not None:
        for b in ppr.findall(qn("w:pBdr")):
            ppr.remove(b)


def _shade(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _borders(table) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), LINE)
        borders.append(el)
    mar = OxmlElement("w:tblCellMar")
    for edge, w in (("top", 60), ("left", 100), ("bottom", 60), ("right", 100)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(w))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    # the schema orders tblPr children: ... tblBorders, shd, tblLayout, tblCellMar, tblLook
    look = tblPr.find(qn("w:tblLook"))
    for el in (borders, mar):
        if look is not None:
            look.addprevious(el)
        else:
            tblPr.append(el)


def _repeat_header(row) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(el)


def _hyperlink(paragraph, url: str, text: str, color: str) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    rid = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(c)
    rpr.append(u)
    run.append(rpr)
    txt = OxmlElement("w:t")
    txt.text = text
    run.append(txt)
    link.append(run)
    paragraph._p.append(link)
